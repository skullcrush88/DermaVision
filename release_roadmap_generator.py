"""
Release Roadmap Generator - Gantt Style

Generates an Excel file with a Gantt-style release plan/roadmap.

REQUIREMENTS:
    - pandas
    - openpyxl

INSTALLATION:
    pip install pandas openpyxl

USAGE:
    python release_roadmap_generator.py

OUTPUT:
    - release_roadmap.xlsx (in the same directory)
"""

import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_gantt_roadmap():
    """Generate a Gantt-style release roadmap Excel file."""
    
    # Define features and their sprint/week assignments
    # Format: 'Feature': [(sprint_num, week_num), ...]
    feature_timeline = {
        'Image Upload': [(1, 1), (1, 2)],
        'Preprocessing': [(1, 2), (1, 3)],
        'AI Prediction': [(1, 3), (1, 4)],
        'Image Validation': [(1, 3), (1, 4)],
        'Severity Detection': [(2, 1), (2, 2)],
        'Recommendations': [(2, 2), (2, 3)],
        'UI Improvements': [(2, 2), (2, 3), (2, 4)],
        'Confidence Logic': [(2, 3), (2, 4)],
        'Edge AI': [(3, 1), (3, 2)],
        'History Tracking': [(3, 1), (3, 2), (3, 3)],
        'Deployment': [(3, 2), (3, 3), (3, 4)],
        'Maintenance': [(1, 1), (1, 2), (1, 3), (1, 4), 
                        (2, 1), (2, 2), (2, 3), (2, 4),
                        (3, 1), (3, 2), (3, 3), (3, 4)]
    }
    
    # Color palette for sprints
    sprint_colors = {
        1: 'B4C7E7',  # Light blue for Sprint 1
        2: 'C6E0B4',  # Light green for Sprint 2
        3: 'F4B084'   # Light orange for Sprint 3
    }
    
    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'Roadmap'
    
    # Set up header row
    ws['A1'] = 'Feature'
    headers = []
    col_idx = 2
    
    for sprint in range(1, 4):
        for week in range(1, 5):
            header = f'S{sprint}W{week}'
            headers.append(header)
            ws.cell(row=1, column=col_idx, value=header)
            col_idx += 1
    
    # Define styles
    bold_font = Font(bold=True, size=11)
    center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Format header row
    for col in range(1, len(headers) + 2):
        cell = ws.cell(row=1, column=col)
        cell.font = bold_font
        cell.alignment = center_alignment
        cell.border = border
    
    # Add features and create Gantt bars
    row = 2
    for feature, timeline in feature_timeline.items():
        ws.cell(row=row, column=1, value=feature)
        
        # Format feature name
        feature_cell = ws.cell(row=row, column=1)
        feature_cell.font = bold_font
        feature_cell.alignment = Alignment(horizontal='left', vertical='center')
        feature_cell.border = border
        
        # Find min and max week for this feature
        if timeline:
            timeline_sorted = sorted(timeline)
            min_sprint, min_week = timeline_sorted[0]
            max_sprint, max_week = timeline_sorted[-1]
            
            # Convert to column positions
            min_col = 2 + (min_sprint - 1) * 4 + (min_week - 1)
            max_col = 2 + (max_sprint - 1) * 4 + (max_week - 1)
            
            # Get the primary sprint color (first sprint in timeline)
            primary_sprint = min_sprint
            bar_color = sprint_colors[primary_sprint]
            bar_fill = PatternFill(start_color=bar_color, end_color=bar_color, fill_type='solid')
            
            # Merge and fill cells for the feature's duration
            if min_col == max_col:
                # Single cell
                cell = ws.cell(row=row, column=min_col, value='')
                cell.fill = bar_fill
                cell.alignment = center_alignment
                cell.border = border
            else:
                # Merge multiple cells
                min_col_letter = get_column_letter(min_col)
                max_col_letter = get_column_letter(max_col)
                ws.merge_cells(f'{min_col_letter}{row}:{max_col_letter}{row}')
                
                merged_cell = ws.cell(row=row, column=min_col)
                merged_cell.value = ''
                merged_cell.fill = bar_fill
                merged_cell.alignment = center_alignment
                merged_cell.border = border
        
        # Format empty cells in the row
        for col in range(2, len(headers) + 2):
            cell = ws.cell(row=row, column=col)
            if cell.value is None or cell.value == '':
                cell.border = border
                cell.alignment = center_alignment
        
        row += 1
    
    # Set column widths
    ws.column_dimensions['A'].width = 20
    for col in range(2, len(headers) + 2):
        ws.column_dimensions[get_column_letter(col)].width = 12
    
    # Set row height for better visibility
    ws.row_dimensions[1].height = 25
    for r in range(2, row):
        ws.row_dimensions[r].height = 30
    
    # Save workbook
    filename = 'release_roadmap.xlsx'
    wb.save(filename)
    print(f"✓ Gantt-style roadmap exported to {filename}")
    print(f"✓ Total features: {len(feature_timeline)}")
    print(f"✓ Timeline: 3 sprints × 4 weeks = 12 weeks")
    print(f"✓ Sprint colors: S1=Blue, S2=Green, S3=Orange")


if __name__ == '__main__':
    generate_gantt_roadmap()
